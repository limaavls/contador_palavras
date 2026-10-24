from collections import Counter
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


def ler_txt(caminho: Path) -> str:
    return caminho.read_text(encoding="utf-8")


def ler_docx(caminho: Path) -> str:
    doc = Document(caminho)
    partes = [paragrafo.text for paragrafo in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                partes.append(celula.text)
    return "\n".join(partes)


def ler_pdf(caminho: Path) -> str:
    leitor = PdfReader(caminho)
    return "\n".join(pagina.extract_text() or "" for pagina in leitor.pages)


LEITORES = {
    ".txt": ler_txt,
    ".docx": ler_docx,
    ".pdf": ler_pdf,
}

arquivo = input("Digite o caminho para o arquivo (.txt, .docx ou .pdf): ")
caminho = Path(arquivo.strip().strip('"').strip("'"))

if not caminho.is_file():
    print("Arquivo não encontrado.")
    raise SystemExit(1)

extensao = caminho.suffix.lower()
leitor = LEITORES.get(extensao)
if leitor is None:
    print(f"Formato {extensao or 'sem extensão'} não suportado. Use .txt, .docx ou .pdf.")
    raise SystemExit(1)

texto = leitor(caminho)
palavras = re.findall(r"\w+", texto.lower())
print(f"O arquivo possui {len(palavras)} palavras.")

palavras_frequentes = Counter(palavras).most_common(10)
print(f"As 10 palavras mais frequentes são: {palavras_frequentes}")
